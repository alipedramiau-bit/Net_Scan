from __future__ import annotations

import sys
from pathlib import Path

# --- Path bootstrap: make this file runnable in every supported way ---
#   python backend/main.py            (from project root)
#   python main.py                    (from inside backend/)
#   python -m backend.main            (from project root)
ROOT_DIR = Path(__file__).resolve().parent.parent
BACKEND_DIR = Path(__file__).resolve().parent
for _p in (str(ROOT_DIR), str(BACKEND_DIR)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from concurrent.futures import ThreadPoolExecutor
from threading import Lock, Event
from datetime import datetime, timezone
from uuid import uuid4

from fastapi import FastAPI, HTTPException, Query, UploadFile, File, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, JSONResponse
from pydantic import BaseModel, Field
import os
import platform
import re as _re
import shutil
import subprocess

# --- Robust imports: package-relative first, plain fallback for direct run ---
try:
    from backend import db
except ImportError:
    import db

from scanner_lib import ScanConfig, ScanReport, scan_targets  # package-level API
from scanner_lib.constants import MAX_TARGETS
from scanner_lib.parser import parse_targets
from scanner_lib.reports import REPORT_KEYS, csv_bytes

# Default target list (CIDRs/IPs) used when the client sends an empty targets field.
# File lives at project root next to backend/ and frontend/.
DEFAULT_TARGETS_FILE = ROOT_DIR / "ip_olom.txt"


def _load_default_targets_text() -> str:
    """Read ip_olom.txt from project root; raise a clear error if missing/empty."""
    path = DEFAULT_TARGETS_FILE
    if not path.is_file():
        raise FileNotFoundError(
            f"فایل پیش‌فرض هدف‌ها یافت نشد: {path.name} "
            f"(مسیر مورد انتظار: {path})"
        )
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        raise ValueError(f"فایل پیش‌فرض هدف‌ها خالی است: {path.name}")
    return text


try:
    from backend import smart_reports
except ImportError:
    import smart_reports

try:
    from backend import vpn_manager
except ImportError:
    import vpn_manager

try:
    from backend import snmp_manager
except ImportError:
    import snmp_manager

try:
    from backend import routeros_api
except ImportError:
    import routeros_api

try:
    from backend import email_manager
except ImportError:
    import email_manager

try:
    from backend import auth
except ImportError:
    import auth

app=FastAPI(title="Authorized IPv4 Scanner")
# allow_origin_regex echoes any origin; combined with allow_credentials=True this
# lets the session cookie work for both same-origin UI and LAN cross-port setups.
app.add_middleware(CORSMiddleware,allow_origin_regex=".*",allow_methods=["*"],allow_headers=["*"],allow_credentials=True)

# Chrome/Edge "Private Network Access": public UI (e.g. http://185.x:5001) calling http://127.0.0.1:8001
# requires Access-Control-Allow-Private-Network on preflight and responses.
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.responses import Response as StarletteResponse

class PrivateNetworkAccessMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        origin = request.headers.get("origin", "*")
        if request.method == "OPTIONS" and request.headers.get("access-control-request-private-network", "").lower() == "true":
            resp = StarletteResponse(status_code=204)
            resp.headers["Access-Control-Allow-Origin"] = origin
            resp.headers["Access-Control-Allow-Methods"] = "GET,POST,PUT,PATCH,DELETE,OPTIONS"
            resp.headers["Access-Control-Allow-Headers"] = request.headers.get("access-control-request-headers", "*")
            resp.headers["Access-Control-Allow-Private-Network"] = "true"
            resp.headers["Access-Control-Max-Age"] = "86400"
            return resp
        response = await call_next(request)
        response.headers["Access-Control-Allow-Private-Network"] = "true"
        if origin and "access-control-allow-origin" not in {k.lower() for k in response.headers.keys()}:
            response.headers["Access-Control-Allow-Origin"] = origin
        return response

app.add_middleware(PrivateNetworkAccessMiddleware)

# ---------------- Admin authentication (session cookie) ----------------
class AuthMiddleware(BaseHTTPMiddleware):
    """Require a valid session cookie for every API route.

    Open paths: login/logout/auth-check (bootstrapping) and everything that is
    not an API prefix (SPA shell, static assets, /health).
    """
    OPEN_EXACT = {"/api/login", "/api/logout", "/api/auth/check"}
    PROTECTED_PREFIXES = ("/api/", "/devices", "/reports", "/scans")

    async def dispatch(self, request, call_next):
        path = request.url.path
        if request.method == "OPTIONS" or path in self.OPEN_EXACT or not path.startswith(self.PROTECTED_PREFIXES):
            return await call_next(request)
        user = auth.verify_token(request.cookies.get(auth.COOKIE_NAME))
        if not user:
            return JSONResponse({"detail": "احراز هویت لازم است"}, status_code=401)
        request.state.user = user
        return await call_next(request)

app.add_middleware(AuthMiddleware)

class LoginBody(BaseModel):
    username: str = Field(default="", max_length=64)
    password: str = Field(default="", max_length=128)


@app.post("/api/login")
def api_login(body: LoginBody, response: Response):
    """Verify admin credentials and set the session cookie."""
    if not auth.check_credentials(body.username, body.password):
        raise HTTPException(401, "نام کاربری یا گذرواژه اشتباه است")
    token = auth.create_token(body.username)
    response.set_cookie(
        auth.COOKIE_NAME, token,
        max_age=auth.COOKIE_MAX_AGE, httponly=True, samesite="lax", path="/",
    )
    return {"ok": True, "username": body.username}


@app.post("/api/logout")
def api_logout(response: Response):
    response.delete_cookie(auth.COOKIE_NAME, path="/")
    return {"ok": True}


@app.get("/api/auth/check")
def api_auth_check(request: Request):
    user = auth.verify_token(request.cookies.get(auth.COOKIE_NAME))
    if not user:
        raise HTTPException(401, "احراز هویت لازم است")
    return {"authenticated": True, "username": user}


JOBS={}; GUARD=Lock(); EXECUTOR=ThreadPoolExecutor(max_workers=4)
PROBE_CANCEL=Event()  # cancel current SNMP probe
API_PROBE_CANCEL=Event()
API_PROBE_PROGRESS={"total":0,"done":0,"current_ip":"","running":False,"canceled":False}
db.init_db_global()

@app.on_event("startup")
def startup():
    db.init_db_global()
def now(): return datetime.now(timezone.utc).isoformat()
class StartRequest(BaseModel):
    # Empty / omitted → load from project-root ip_olom.txt
    targets: str = Field(default="", description="لیست IP/CIDR؛ خالی = استفاده از ip_olom.txt")
    ports: list[int] = Field(default_factory=lambda: [8291, 8728], min_length=1, max_length=128)
    workers: int = Field(default=32, ge=1, le=256)
    ping_timeout_ms: int = Field(default=1000, ge=100, le=10000)
    port_timeout: float = Field(default=1.0, ge=.1, le=10)
    max_targets: int = Field(default=MAX_TARGETS, ge=1, le=MAX_TARGETS)
def public(job):
    report=job.get("report")
    return {"scan_id":job["id"],"status":job["status"],"started_at":job.get("started_at"),"ended_at":job.get("ended_at"),"progress":job["progress"],"records":[{"ip":x.ip,"classification":x.classification,"ping_responded":x.ping_responded,"open_ports":[{"port":p,"state":s} for p,s in x.ports.items() if s=="open"],"states":{"8291":x.ports.get(8291,"not_scanned"),"8728":x.ports.get(8728,"not_scanned")},"port_states":x.ports,"error":x.error} for x in report.records] if report else [],"summary":{k:sum(1 for x in report.records if predicate(x)) for k,predicate in [("ping_responders",lambda x:x.is_alive),("mikrotik_candidates",lambda x:x.classification=="mikrotik_candidate"),("potential_non_mikrotik",lambda x:x.classification=="potential_non_mikrotik"),("winbox_open_api_closed",lambda x:x.classification=="winbox_open_api_closed")]} if report else {}}
def run(job,targets,config):
    try:
        with GUARD:
            if job["cancel_event"].is_set():
                job["status"]="canceled";job["ended_at"]=now();return
            job["status"]="running";job["started_at"]=now()
        def progress(p):
            with GUARD: job["progress"]=dict(p)
        report=scan_targets(targets,config,progress,job["cancel_event"])
        # Attach the report and live summary BEFORE persistence so the UI cards
        # are populated at 100% progress even if DB work is slow or fails.
        # Persist hosts that answered ICMP OR have any open TCP port
        # (MikroTik often blocks ICMP while leaving Winbox/API open).
        hosts=[{"ip":h.ip,"ping_ok":h.ping_responded,"open_ports":list(h.open_ports),
                "port_states":{str(port):state for port,state in h.ports.items()},
                "classification":h.classification}
               for h in report.records if h.is_alive]
        with GUARD:
            job["report"]=report
            job["progress"]={"completed":report.completed,"total":report.total_targets,"ping_responders":report.ping_responders}
        # Persist every live host in ONE batched transaction (fast, single commit).
        try:
            db.upsert_devices_batch(hosts)
        except Exception as db_err:  # persistence failure must not hang the job
            with GUARD: job.setdefault("warnings",[]).append(f"db persist failed: {db_err}")
        finally:
            with GUARD:
                job["status"]="canceled" if job["cancel_event"].is_set() else "done"
                job["ended_at"]=now()
    except Exception as e:
        # Robust terminal transition: status ALWAYS becomes error (never stuck on running).
        try:
            with GUARD:
                job["status"]="error";job["error"]=str(e);job["ended_at"]=now()
        except Exception:
            pass
    finally:
        if job.get("status") not in ("done","error","canceled"):
            with GUARD:
                job["status"]="error";job.setdefault("error","unexpected scan termination");job["ended_at"]=now()
@app.get("/health")
def health(): return {"status":"ok", "database": db.get_active_db_name()}


# ---- multi-database management ----
class DatabaseNameBody(BaseModel):
    name: str = Field(min_length=1, max_length=64)


@app.get("/api/databases")
def api_list_databases():
    return db.list_databases()


@app.post("/api/databases")
def api_create_database(body: DatabaseNameBody):
    try:
        return db.create_database(body.name)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@app.put("/api/databases/active")
def api_set_active_database(body: DatabaseNameBody):
    try:
        return db.set_active_database(body.name)
    except KeyError:
        raise HTTPException(404, "دیتابیس یافت نشد") from None
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@app.delete("/api/databases/{name}")
def api_delete_database(name: str):
    try:
        return db.delete_database(name)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


class DedupeBody(BaseModel):
    by: str = Field(default="device_name")


@app.post("/devices/dedupe")
def api_dedupe_devices(body: DedupeBody = DedupeBody()):
    try:
        return db.dedupe_devices(by=body.by)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


class DeleteDevicesBody(BaseModel):
    ips: list[str] = Field(min_length=1, max_length=5000)


@app.post("/devices/delete")
def api_delete_devices(body: DeleteDevicesBody):
    return db.delete_devices(body.ips)


class TransferBody(BaseModel):
    database: str = Field(min_length=1, max_length=64)


@app.post("/scans/{scan_id}/transfer")
def transfer_scan_to_db(scan_id: str, body: TransferBody):
    with GUARD:
        job = JOBS.get(scan_id)
        if not job:
            raise HTTPException(404, "شناسه اسکن یافت نشد")
        report = job.get("report")
        if not report or job.get("status") not in ("done", "completed", "canceled"):
            raise HTTPException(409, "گزارش اسکن هنوز آماده نیست")
        hosts = [
            {
                "ip": h.ip,
                "ping_ok": h.ping_responded,
                "open_ports": list(h.open_ports),
                "port_states": {str(port): state for port, state in h.ports.items()},
                "classification": h.classification,
            }
            for h in report.records
            if h.is_alive
        ]
    try:
        result = db.transfer_hosts_to_db(hosts, body.database)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    return result


@app.get("/devices")
def get_devices(
    classification: str | None = Query(
        default=None,
        pattern="^(mikrotik_candidate|potential_non_mikrotik|winbox_open_api_closed|non_mikrotik|special)$",
    ),
    device_type: str | None = Query(default=None, pattern="^(AP|ST|Router)$"),
    has_port: int | None = Query(default=None, ge=1, le=65535),
    q: str | None = Query(default=None, max_length=128),
    q_fields: str | None = Query(default=None, max_length=256,
                                 description="محدوده جستجو: شناسه فیلدها جدا شده با کاما (خالی = همه فیلدها)"),
    ping_ok: bool | None = Query(default=None),
    new_only: bool = Query(default=False, description="فقط دستگاه‌های جدید"),
    sort: str = Query(default="ip", pattern="^(ip|ping|ports|ports_count|classification|device_type|first_seen|last_seen|seen_count|device_name|tags)$"),
    order: str = Query(default="asc", pattern="^(asc|desc)$"),
    limit: int = Query(default=500, ge=1, le=200000),
    offset: int = Query(default=0, ge=0),
):
    return db.query_devices(
        classification=classification,
        has_port=has_port,
        q=q,
        q_fields=q_fields,
        limit=limit,
        offset=offset,
        ping_ok=ping_ok,
        new_only=new_only,
        sort=sort,
        order=order,
        device_type=device_type,
    )

@app.get("/devices/export.csv")
def export_devices(
    classification: str | None = Query(
        default=None,
        pattern="^(mikrotik_candidate|potential_non_mikrotik|winbox_open_api_closed|non_mikrotik|special)$",
    ),
    has_port: int | None = Query(default=None, ge=1, le=65535),
    q: str | None = Query(default=None, max_length=128),
    q_fields: str | None = Query(default=None, max_length=256,
                                 description="محدوده جستجو: شناسه فیلدها جدا شده با کاما (خالی = همه فیلدها)"),
    ping_ok: bool | None = Query(default=None),
    new_only: bool = Query(default=False, description="فقط دستگاه‌های جدید"),
    sort: str = Query(default="ip", pattern="^(ip|ping|ports|ports_count|classification|first_seen|last_seen|seen_count|device_name|tags)$"),
    order: str = Query(default="asc", pattern="^(asc|desc)$"),
    limit: int = Query(default=10000, ge=1, le=200000),
    offset: int = Query(default=0, ge=0),
):
    result = db.query_devices(classification, has_port, q, limit, offset, ping_ok=ping_ok,
                              new_only=new_only, sort=sort, order=order, q_fields=q_fields)
    filename = "new_devices.csv" if new_only else "devices.csv"
    return Response(db.devices_to_csv(result["devices"]), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


class ExportSelectedBody(BaseModel):
    ips: list[str] = Field(default_factory=list, max_length=50000)


@app.post("/devices/export.xlsx")
def export_devices_xlsx(body: ExportSelectedBody):
    """Export selected device records as an Excel .xlsx file for download."""
    ips = [str(x).strip() for x in (body.ips or []) if str(x).strip()]
    if not ips:
        raise HTTPException(400, "هیچ دستگاهی انتخاب نشده است")
    devices = db.fetch_devices_by_ips(ips)
    if not devices:
        raise HTTPException(404, "دستگاه انتخاب‌شده در دیتابیس فعال یافت نشد")
    data = db.devices_to_xlsx(devices)
    from datetime import datetime, timezone
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"devices_selected_{stamp}.xlsx"
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/devices/new/count")
def new_devices_count():
    """Summary metric: دستگاه‌های جدید اضافه‌شده در این اسکن/دوره."""
    result = db.query_devices(new_only=True, limit=1)
    return {"new_count": result["new_count"]}

class ManualDeviceRequest(BaseModel):
    ip: str = Field(max_length=45)
    open_ports: list[int] = Field(default_factory=list)
    classification: str = Field(default="mikrotik_candidate", max_length=64)

@app.post("/devices/manual")
def add_manual_device(body: ManualDeviceRequest):
    try:
        return db.add_manual_device(body.ip.strip(), body.open_ports, body.classification)
    except ValueError as e:
        raise HTTPException(422, f"ورودی نامعتبر: {e}")


# ---- backup / restore devices.db ----
@app.get("/devices/backup")
def backup_devices_db(label: str | None = Query(default=None, max_length=64)):
    """Download a consistent snapshot of the *active* database as a separate backup file."""
    try:
        data = db.backup_db_bytes()
    except FileNotFoundError as e:
        raise HTTPException(404, str(e)) from e
    except Exception as e:
        raise HTTPException(500, f"پشتیبان‌گیری ناموفق: {e}") from e
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    dbname = db.get_active_db_name()
    safe_label = ""
    if label:
        safe_label = "_" + "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in label)[:40]
    filename = f"backup_{dbname}{safe_label}_{ts}.db"
    return Response(
        content=data,
        media_type="application/x-sqlite3",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Backup-Database": dbname,
            "X-Backup-Filename": filename,
        },
    )


@app.post("/devices/restore")
async def restore_devices_db(file: UploadFile = File(...)):
    """Upload a backup .db and replace the *active* database (previous copy kept as .pre_restore)."""
    if not file.filename:
        raise HTTPException(400, "فایلی انتخاب نشده است")
    if not str(file.filename).lower().endswith(".db"):
        # still allow if content is sqlite
        pass
    data = await file.read()
    try:
        result = db.restore_db_bytes(data)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    except Exception as e:
        raise HTTPException(500, f"بازیابی ناموفق: {e}") from e
    result["active_database"] = db.get_active_db_name()
    result["source_filename"] = file.filename
    return result


# ---- import tags from Excel ----
@app.post("/devices/import-tags")
async def import_tags_xlsx(file: UploadFile = File(...)):
    """Import tags from Excel (.xlsx) with columns IP and TAG; update matching devices."""
    if not file.filename:
        raise HTTPException(400, "فایلی انتخاب نشده است")
    name = str(file.filename).lower()
    if not (name.endswith(".xlsx") or name.endswith(".xlsm") or name.endswith(".xls")):
        raise HTTPException(400, "فقط فایل Excel (.xlsx) پذیرفته می‌شود")
    data = await file.read()
    if not data:
        raise HTTPException(400, "فایل خالی است")
    try:
        from openpyxl import load_workbook
        import io as _io
        wb = load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise HTTPException(400, "فایل Excel خالی است")
        # map header names (case-insensitive, strip)
        def norm(v):
            return str(v or "").strip().lower().replace(" ", "").replace("_", "")
        headers = [norm(h) for h in header]
        ip_idx = None
        tag_idx = None
        for i, h in enumerate(headers):
            if h in ("ip", "address", "host", "ایپی", "آدرس"):
                ip_idx = i
            if h in ("tag", "tags", "تگ", "برچسب"):
                tag_idx = i
        if ip_idx is None:
            # default first column = IP, second = TAG
            ip_idx = 0
        if tag_idx is None:
            tag_idx = 1 if len(headers) > 1 else 0
        items = []
        for row in rows_iter:
            if not row:
                continue
            ip_val = row[ip_idx] if ip_idx < len(row) else None
            tag_val = row[tag_idx] if tag_idx < len(row) else ""
            if ip_val is None or str(ip_val).strip() == "":
                continue
            items.append({"ip": str(ip_val).strip(), "tags": "" if tag_val is None else str(tag_val).strip()})
        wb.close()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"خواندن Excel ناموفق: {e}") from e
    if not items:
        raise HTTPException(400, "هیچ ردیف IP/TAG در فایل یافت نشد")
    result = db.apply_tags_by_ip(items)
    result["source_filename"] = file.filename
    return result


# ---- full project backup / restore (ZIP) ----
@app.get("/api/backup/full")
def full_backup_zip():
    """Download a ZIP containing all databases and config files."""
    try:
        data = db.full_backup_zip_bytes()
    except Exception as e:
        raise HTTPException(500, f"پشتیبان‌گیری کامل ناموفق: {e}") from e
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"mikrotik_scanner_full_{ts}.zip"
    return Response(
        content=data,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Backup-Filename": filename,
        },
    )


@app.post("/api/backup/full/restore")
async def full_restore_zip(file: UploadFile = File(...)):
    """Restore all databases and configs from a full-backup ZIP."""
    if not file.filename:
        raise HTTPException(400, "فایلی انتخاب نشده است")
    data = await file.read()
    try:
        result = db.full_restore_from_zip(data)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    except Exception as e:
        raise HTTPException(500, f"بازیابی کامل ناموفق: {e}") from e
    result["source_filename"] = file.filename
    return result


# ---- settings (generic key/value) ----
class SettingBody(BaseModel):
    value: str = Field(default="", max_length=8192)


@app.get("/api/settings/{key}")
def get_setting(key: str):
    if not _re.fullmatch(r"[a-zA-Z0-9_.-]{1,64}", key or ""):
        raise HTTPException(400, "کلید تنظیمات نامعتبر است")
    return {"key": key, "value": db.get_setting(key)}


@app.put("/api/settings/{key}")
def put_setting(key: str, body: SettingBody):
    if not _re.fullmatch(r"[a-zA-Z0-9_.-]{1,64}", key or ""):
        raise HTTPException(400, "کلید تنظیمات نامعتبر است")
    return db.set_setting(key, body.value)


# ---- multi Winbox path management ----
# Stored in settings as JSON list under key "winbox_list":
# [{"id":"wb1","name":"Winbox64","path":"D:\\winbox64\\winbox64.exe","is_default":true}, ...]
import json as _json
import uuid as _uuid


def _winbox_tools_dir() -> Path:
    return ROOT_DIR / "tools" / "winbox"


def _scan_bundled_winbox() -> list[dict]:
    """Discover .exe files placed under tools/winbox/."""
    found = []
    d = _winbox_tools_dir()
    if not d.is_dir():
        return found
    for p in sorted(d.iterdir()):
        if p.is_file() and p.suffix.lower() in (".exe", "") and p.name.lower() not in ("readme.txt", "readme.md"):
            if p.suffix.lower() == ".exe" or p.name.lower().startswith("winbox"):
                found.append(
                    {
                        "id": "bundled-" + _re.sub(r"[^a-zA-Z0-9_-]", "-", p.stem)[:32],
                        "name": p.name,
                        "path": str(p.resolve()),
                        "is_default": False,
                        "bundled": True,
                        "exists": True,
                    }
                )
    return found


def _load_winbox_list() -> list[dict]:
    raw = db.get_setting("winbox_list")
    items: list[dict] = []
    if raw:
        try:
            parsed = _json.loads(raw)
            if isinstance(parsed, list):
                for it in parsed:
                    if not isinstance(it, dict):
                        continue
                    path = str(it.get("path") or "").strip()
                    if not path:
                        continue
                    items.append(
                        {
                            "id": str(it.get("id") or _uuid.uuid4().hex[:10]),
                            "name": str(it.get("name") or Path(path).name)[:64],
                            "path": path,
                            "is_default": bool(it.get("is_default")),
                            "bundled": bool(it.get("bundled")),
                            "exists": Path(path).is_file(),
                        }
                    )
        except (ValueError, TypeError):
            items = []
    # migrate legacy single path
    legacy = db.get_setting("winbox_custom_path")
    if legacy and not any(i["path"].lower() == legacy.lower() for i in items):
        items.insert(
            0,
            {
                "id": "legacy",
                "name": Path(legacy).name or "Winbox",
                "path": legacy,
                "is_default": not any(i.get("is_default") for i in items),
                "bundled": False,
                "exists": Path(legacy).is_file(),
            },
        )
    # merge discovered bundled tools (do not duplicate paths)
    for b in _scan_bundled_winbox():
        if not any(i["path"].lower() == b["path"].lower() for i in items):
            items.append(b)
    if items and not any(i.get("is_default") for i in items):
        # prefer an existing file as default
        for i in items:
            if i.get("exists"):
                i["is_default"] = True
                break
        else:
            items[0]["is_default"] = True
    return items


def _save_winbox_list(items: list[dict]) -> list[dict]:
    # persist without runtime-only fields
    payload = [
        {
            "id": i["id"],
            "name": i["name"],
            "path": i["path"],
            "is_default": bool(i.get("is_default")),
            "bundled": bool(i.get("bundled")),
        }
        for i in items
    ]
    db.set_setting("winbox_list", _json.dumps(payload, ensure_ascii=False))
    # keep legacy key in sync with default
    default = next((i for i in items if i.get("is_default")), items[0] if items else None)
    if default:
        db.set_setting("winbox_custom_path", default["path"])
    return _load_winbox_list()


class WinboxPathBody(BaseModel):
    name: str = Field(default="", max_length=64)
    path: str = Field(min_length=1, max_length=512)
    is_default: bool = False


@app.get("/api/winbox/list")
def winbox_list():
    return {"items": _load_winbox_list(), "tools_dir": str(_winbox_tools_dir())}


@app.post("/api/winbox/list")
def winbox_add(body: WinboxPathBody):
    path = body.path.strip().strip('"')
    items = _load_winbox_list()
    # update if same path exists
    existing = next((i for i in items if i["path"].lower() == path.lower()), None)
    if existing:
        existing["name"] = (body.name or existing["name"])[:64]
        if body.is_default:
            for i in items:
                i["is_default"] = i["id"] == existing["id"]
        return {"items": _save_winbox_list(items), "id": existing["id"]}
    new_id = "wb-" + _uuid.uuid4().hex[:10]
    entry = {
        "id": new_id,
        "name": (body.name or Path(path).name or "Winbox")[:64],
        "path": path,
        "is_default": bool(body.is_default) or not items,
        "bundled": False,
        "exists": Path(path).is_file(),
    }
    if entry["is_default"]:
        for i in items:
            i["is_default"] = False
    items.append(entry)
    return {"items": _save_winbox_list(items), "id": new_id}


@app.put("/api/winbox/list/{item_id}/default")
def winbox_set_default(item_id: str):
    items = _load_winbox_list()
    if not any(i["id"] == item_id for i in items):
        raise HTTPException(404, "مسیر Winbox یافت نشد")
    for i in items:
        i["is_default"] = i["id"] == item_id
    return {"items": _save_winbox_list(items)}


@app.delete("/api/winbox/list/{item_id}")
def winbox_delete(item_id: str):
    items = _load_winbox_list()
    new_items = [i for i in items if i["id"] != item_id]
    if len(new_items) == len(items):
        raise HTTPException(404, "مسیر Winbox یافت نشد")
    if new_items and not any(i.get("is_default") for i in new_items):
        new_items[0]["is_default"] = True
    return {"items": _save_winbox_list(new_items)}


def _resolve_winbox_exe(preferred_id: str | None = None) -> str | None:
    items = _load_winbox_list()
    if preferred_id:
        for i in items:
            if i["id"] == preferred_id and Path(i["path"]).is_file():
                return i["path"]
    for i in items:
        if i.get("is_default") and Path(i["path"]).is_file():
            return i["path"]
    for i in items:
        if Path(i["path"]).is_file():
            return i["path"]
    # auto-discover common locations
    if platform.system().lower() == "windows":
        candidates = [
            shutil.which("winbox"),
            shutil.which("winbox64"),
            shutil.which("winbox.exe"),
            shutil.which("winbox64.exe"),
            r"C:\Program Files\WinBox\winbox64.exe",
            r"C:\Program Files\WinBox\winbox.exe",
            r"C:\Program Files (x86)\WinBox\winbox.exe",
            str(Path.home() / "Desktop" / "winbox64.exe"),
            str(Path.home() / "Desktop" / "winbox.exe"),
            str(Path.home() / "Downloads" / "winbox64.exe"),
            str(Path.home() / "Downloads" / "winbox.exe"),
        ]
        for c in candidates:
            if c and Path(c).is_file():
                return c
    else:
        for name in ("winbox", "wine"):
            p = shutil.which(name)
            if p:
                return p
    return None


@app.post("/devices/{ip}/winbox")
def open_winbox(ip: str, winbox_id: str | None = Query(default=None)):
    """Launch Winbox against the given IP (Windows host running the backend)."""
    try:
        parts = ip.strip().split(".")
        if len(parts) != 4 or any(not p.isdigit() or not 0 <= int(p) <= 255 for p in parts):
            raise ValueError
    except ValueError:
        raise HTTPException(400, "آدرس IP نامعتبر است") from None
    exe = _resolve_winbox_exe(winbox_id)
    if not exe:
        raise HTTPException(
            404,
            "Winbox پیدا نشد. از بخش «مدیریت Winbox» یک مسیر اضافه کنید "
            f"یا فایل را در پوشه {_winbox_tools_dir()} قرار دهید.",
        )
    try:
        if Path(exe).name.lower().startswith("wine"):
            subprocess.Popen([exe, "winbox.exe", ip], close_fds=True)
        else:
            subprocess.Popen([exe, ip], close_fds=True, cwd=str(Path(exe).parent))
    except OSError as e:
        raise HTTPException(500, f"اجرای Winbox ناموفق: {e}") from e
    return {"ok": True, "ip": ip, "exe": exe}


@app.get("/devices/{ip}")
def get_device(ip: str):
    result = db.query_devices(q=ip, limit=10000)
    for device in result["devices"]:
        if device["ip"] == ip:
            return device
    raise HTTPException(404, "دستگاه یافت نشد")

@app.post("/scans", status_code=202)
def start(body: StartRequest):
    raw = (body.targets or "").strip()
    used_default = False
    if not raw:
        try:
            raw = _load_default_targets_text()
            used_default = True
        except (OSError, ValueError) as e:
            raise HTTPException(422, str(e)) from e
    try:
        targets = parse_targets(raw, body.max_targets)
    except ValueError as e:
        raise HTTPException(422, str(e)) from e
    ports = tuple(dict.fromkeys(body.ports + [8291, 8728]))
    if any(not 1 <= p <= 65535 for p in ports):
        raise HTTPException(422, "پورت باید بین ۱ تا ۶۵۵۳۵ باشد")
    sid = str(uuid4())
    job = {
        "id": sid,
        "status": "queued",
        "started_at": None,
        "ended_at": None,
        "cancel_event": Event(),
        "progress": {"completed": 0, "total": len(targets), "ping_responders": 0},
        "report": None,
        "used_default_targets": used_default,
    }
    with GUARD:
        JOBS[sid] = job
    EXECUTOR.submit(
        run,
        job,
        targets,
        ScanConfig(ports, body.workers, body.ping_timeout_ms, body.port_timeout, body.max_targets),
    )
    return {
        "scan_id": sid,
        "status": "queued",
        "total_targets": len(targets),
        "used_default_targets": used_default,
    }
@app.post("/scans/{scan_id}/cancel")
def cancel(scan_id:str):
    with GUARD:
        job=JOBS.get(scan_id)
        if not job: raise HTTPException(404,"شناسه اسکن یافت نشد")
        if job["status"] in ("queued","running"):
            job["cancel_event"].set()
            if job["status"]=="queued":job["status"]="canceled";job["ended_at"]=now()
        return {"scan_id":scan_id,"status":job["status"],"cancel_requested":job["cancel_event"].is_set()}
@app.get("/scans/{scan_id}")
def status(scan_id:str):
    job=JOBS.get(scan_id)
    if not job: raise HTTPException(404,"شناسه اسکن یافت نشد")
    return public(job)|({"error":job["error"]} if job.get("error") else {})
@app.get("/scans/{scan_id}/reports/{key}")
def download(scan_id:str,key:str):
    job=JOBS.get(scan_id)
    if not job: raise HTTPException(404,"شناسه اسکن یافت نشد")
    if key not in REPORT_KEYS: raise HTTPException(404,"گزارش ناشناخته")
    if job.get("status") not in ("done","canceled"): raise HTTPException(409,"گزارش هنوز آماده نیست")
    return Response(csv_bytes(job["report"],key),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":f'attachment; filename="{key}.csv"'})


# ---------------- Smart reports: interface flaps ----------------
def _resolve_report_database(database: str = "") -> str:
    """Validate a database name for smart reports; empty means the active DB."""
    name = (database or "").strip()
    if not name:
        return db.get_active_db_name()
    try:
        known = {item["name"] for item in db.list_databases()["databases"]}
    except Exception:
        known = set()
    if name not in known:
        raise HTTPException(404, f"دیتابیس «{name}» یافت نشد")
    return name


@app.get("/reports/interface-flaps")
def report_interface_flaps(days: int = 30, threshold_minutes: int = 5, database: str = Query(default="", max_length=64)):
    db_name = _resolve_report_database(database)
    error_msg = None
    try:
        devices = db.query_devices(limit=100000, db_path=str(db.db_path_for(db_name)))
    except Exception as exc:
        devices = []
        error_msg = f"خواندن دستگاه‌ها از دیتابیس «{db_name}» ناموفق بود: {exc}"
    rows = devices.get("devices") if isinstance(devices, dict) else devices
    report = smart_reports.interface_flap_report(rows or [], days_back=days, threshold_minutes=threshold_minutes)
    report["database"] = db_name
    if error_msg:
        report["error"] = error_msg
    # Keep the report metadata and provide stable names for both API clients and the UI.
    items = []
    for event in report.get("events", []):
        item = dict(event)
        item.setdefault("iface", item.get("interface", ""))
        item.setdefault("last_flap_time", item.get("datetime", ""))
        item.setdefault("last_event", item.get("last_status", ""))
        items.append(item)
    report["events"] = items
    report["items"] = items
    return report

@app.get("/reports/interface-flaps/export.csv")
def report_interface_flaps_csv(days: int = 30, threshold_minutes: int = 5, database: str = Query(default="", max_length=64)):
    db_name = _resolve_report_database(database)
    try:
        fetched = db.query_devices(limit=100000, db_path=str(db.db_path_for(db_name)))
    except Exception as exc:
        raise HTTPException(500, f"خواندن دستگاه‌ها از دیتابیس «{db_name}» ناموفق بود: {exc}")
    rows = fetched.get("devices", []) if isinstance(fetched, dict) else fetched
    rep = smart_reports.interface_flap_report(rows or [], days_back=days, threshold_minutes=threshold_minutes)
    rep["database"] = db_name
    csv_content = smart_reports.interface_flaps_csv(rep)
    return Response(content=csv_content, media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="interface_flaps_report.csv"'})


@app.get("/reports/interface-flaps/export.xlsx")
def report_interface_flaps_xlsx(days: int = 30, threshold_minutes: int = 5, database: str = Query(default="", max_length=64)):
    db_name = _resolve_report_database(database)
    try:
        devices = db.query_devices(limit=100000, db_path=str(db.db_path_for(db_name)))
    except Exception as exc:
        raise HTTPException(500, f"خواندن دستگاه‌ها از دیتابیس «{db_name}» ناموفق بود: {exc}")
    rows = devices.get("devices") if isinstance(devices, dict) else devices
    rep = smart_reports.interface_flap_report(rows or [], days_back=days, threshold_minutes=threshold_minutes)
    rep["database"] = db_name
    data = smart_reports.interface_flaps_xlsx(rep)
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="interface_flaps.xlsx"'},
    )

# ---------------- Smart report: wireless antenna / station issues ----------------
@app.get("/reports/wireless-antenna-issues")
def report_wireless_antenna_issues(days: int = 30, database: str = Query(default="", max_length=64)):
    db_name = _resolve_report_database(database)
    error_msg = None
    try:
        result = db.query_devices(limit=100000, db_path=str(db.db_path_for(db_name)))
        rows = result.get("devices", []) if isinstance(result, dict) else result
    except Exception as exc:
        rows = []
        error_msg = f"خواندن دستگاه‌ها از دیتابیس «{db_name}» ناموفق بود: {exc}"
    report = smart_reports.wireless_antenna_report(rows or [], days_back=days)
    report["database"] = db_name
    if error_msg:
        report["error"] = error_msg
    return report


@app.get("/reports/wireless-antenna-issues/export.csv")
def report_wireless_antenna_issues_csv(days: int = 30, database: str = Query(default="", max_length=64)):
    db_name = _resolve_report_database(database)
    try:
        result = db.query_devices(limit=100000, db_path=str(db.db_path_for(db_name)))
    except Exception as exc:
        raise HTTPException(500, f"خواندن دستگاه‌ها از دیتابیس «{db_name}» ناموفق بود: {exc}")
    rows = result.get("devices", []) if isinstance(result, dict) else result
    report = smart_reports.wireless_antenna_report(rows or [], days_back=days)
    return Response(
        content=smart_reports.wireless_antenna_csv(report),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="wireless_antenna_issues.csv"'},
    )

# ---------------- VPN connection manager ----------------
class VpnConnectRequest(BaseModel):
    protocol: str = Field(pattern="^(wireguard|openvpn|pptp|l2tp|sstp)$")
    name: str = Field(default="Remote VPN", max_length=64)
    host: str | None = Field(default=None, max_length=253)
    username: str | None = Field(default=None, max_length=128)
    password: str | None = Field(default=None, max_length=128)
    profile_name: str | None = Field(default=None, max_length=128)
    config_content: str | None = Field(default=None, max_length=200_000)
    routes: list[str] | str | None = None

class VpnProfileRequest(BaseModel):
    id: str = Field(pattern="^[a-zA-Z0-9_-]{1,48}$")
    name: str = Field(default="", max_length=64)
    protocol: str = Field(pattern="^(wireguard|openvpn|pptp|l2tp|sstp)$")
    host: str | None = Field(default=None, max_length=253)
    port: int | str | None = Field(default=None)
    username: str | None = Field(default=None, max_length=128)
    profile_name: str | None = Field(default=None, max_length=128)
    config_content: str | None = Field(default=None, max_length=200_000)
    routes: list[str] | str | None = None
    is_default: bool = False

@app.get("/api/vpn/status")
def vpn_status():
    return vpn_manager.status()

@app.post("/api/vpn/connect")
def vpn_connect(body: VpnConnectRequest):
    try:
        return vpn_manager.connect(body.model_dump())
    except (ValueError, RuntimeError) as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise HTTPException(500, f"خطای داخلی VPN: {e}")

@app.post("/api/vpn/disconnect")
def vpn_disconnect():
    try:
        return vpn_manager.disconnect()
    except Exception as e:
        raise HTTPException(500, f"خطا در قطع اتصال VPN: {e}")

@app.get("/api/vpn/profiles")
def vpn_profiles_list():
    return {"profiles": vpn_manager.list_profiles()}

@app.post("/api/vpn/profiles")
def vpn_profiles_save(body: VpnProfileRequest):
    try:
        return vpn_manager.save_profile(body.model_dump())
    except ValueError as e:
        raise HTTPException(400, str(e))

@app.delete("/api/vpn/profiles/{profile_id}")
def vpn_profiles_delete(profile_id: str):
    if vpn_manager.delete_profile(profile_id):
        return {"deleted": profile_id}
    raise HTTPException(404, "پروفایل یافت نشد")


# ---------------- SNMP community profiles (MikroTik-style) ----------------
class SnmpProfileRequest(BaseModel):
    id: str | None = Field(default=None, max_length=64)
    name: str | None = Field(default=None, max_length=64)
    community: str | None = Field(default=None, max_length=64)
    addresses: str | None = Field(default="0.0.0.0/0", max_length=128)
    write_access: bool = False
    version: str = Field(default="2c", max_length=8)
    port: int = Field(default=161, ge=1, le=65535)
    timeout: float = Field(default=2.0, ge=0.2, le=30)
    is_default: bool = False
    raw: str | None = Field(default=None, max_length=4000)


class SnmpProbeRequest(BaseModel):
    ips: list[str] = Field(min_length=1, max_length=1000)
    profile_id: str | None = None
    save_names: bool = True
    api_username: str | None = Field(default=None, max_length=128)
    api_password: str | None = Field(default=None, max_length=128)
    api_port: int = Field(default=8728, ge=1, le=65535)


@app.get("/api/snmp/profiles")
def snmp_profiles_list():
    return {"profiles": snmp_manager.list_profiles()}


@app.post("/api/snmp/profiles")
def snmp_profiles_save(body: SnmpProfileRequest):
    try:
        return snmp_manager.save_profile(body.model_dump())
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@app.put("/api/snmp/profiles/{profile_id}/default")
def snmp_profiles_default(profile_id: str):
    try:
        return snmp_manager.set_default(profile_id)
    except KeyError:
        raise HTTPException(404, "پروفایل SNMP یافت نشد") from None


@app.delete("/api/snmp/profiles/{profile_id}")
def snmp_profiles_delete(profile_id: str):
    if snmp_manager.delete_profile(profile_id):
        return {"deleted": profile_id}
    raise HTTPException(404, "پروفایل SNMP یافت نشد")


class SmtpConfigRequest(BaseModel):
    enabled: bool | None = None
    mode: str | None = Field(default=None, max_length=16)  # direct | relay
    host: str | None = Field(default=None, max_length=256)
    helo_host: str | None = Field(default=None, max_length=256)
    port: int | None = Field(default=None, ge=1, le=65535)
    use_tls: bool | None = None
    use_ssl: bool | None = None
    username: str | None = Field(default=None, max_length=256)
    password: str | None = Field(default=None, max_length=256)
    from_email: str | None = Field(default=None, max_length=256)
    from_name: str | None = Field(default=None, max_length=128)
    to_email: str | None = Field(default=None, max_length=256)
    timeout: float | None = Field(default=None, ge=5, le=120)


class SmtpTestRequest(BaseModel):
    to_email: str | None = Field(default=None, max_length=256)
    message: str | None = Field(default=None, max_length=5000)


@app.get("/api/smtp/config")
def smtp_config_get():
    return email_manager.get_config(mask_password=True)


@app.put("/api/smtp/config")
def smtp_config_save(body: SmtpConfigRequest):
    try:
        data = body.model_dump(exclude_none=True) if hasattr(body, "model_dump") else body.dict(exclude_none=True)
        return email_manager.save_config(data)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    except Exception as e:
        raise HTTPException(500, f"خطا در ذخیره SMTP: {e}") from e


@app.post("/api/smtp/test")
def smtp_test(body: SmtpTestRequest):
    """Send a test email via configured SMTP (default Gmail)."""
    try:
        return email_manager.send_test(to_email=body.to_email or None, message=body.message or None)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    except Exception as e:
        raise HTTPException(502, f"ارسال ناموفق: {e}") from e



@app.post("/api/snmp/probe")
def snmp_probe(body: SnmpProbeRequest):
    """Probe selected devices via SNMP. Supports cancellation via /api/snmp/probe/cancel."""
    PROBE_CANCEL.clear()
    try:
        results = snmp_manager.probe_names(body.ips, body.profile_id, cancel_event=PROBE_CANCEL, api_username=body.api_username, api_password=body.api_password, api_port=body.api_port)
    except RuntimeError as e:
        raise HTTPException(400, str(e)) from e
    canceled = PROBE_CANCEL.is_set()
    saved = 0
    active_db = db.get_active_db_name()
    from pathlib import Path as _P
    txt_dir = _P(db.db_path_for()).parent / "ip_info"
    txt_dir.mkdir(parents=True, exist_ok=True)
    if True:  # always persist probe results to DB
        for r in results:
            if not r.get("ok"):
                continue
            ip = r["ip"]
            name = r.get("device_name") or ""
            interfaces = r.get("interfaces") or ""
            ip_addrs = r.get("ip_addresses") or ""
            device_type = r.get("device_type") or ""
            wireless_info = r.get("wireless_info") or ""
            try:
                kwargs = {}
                if name:
                    kwargs["device_name"] = name
                if (interfaces or "").strip():
                    kwargs["interfaces"] = interfaces
                if (ip_addrs or "").strip():
                    kwargs["ip_addresses"] = ip_addrs
                if device_type:
                    kwargs["device_type"] = device_type
                if (wireless_info or "").strip():
                    kwargs["wireless_info"] = wireless_info
                neigh = r.get("ip_neighbors") or ""
                if neigh.strip():
                    kwargs["ip_neighbors"] = neigh
                if not kwargs:
                    continue
                try:
                    db.update_device_fields(ip, **kwargs)
                except KeyError:
                    db.add_manual_device(ip, open_ports=[161], classification="mikrotik_candidate")
                    db.update_device_fields(ip, **kwargs)
                saved += 1
                safe_ip = ip.replace(":", "_").replace("/", "_")
                fpath = txt_dir / f"{safe_ip}.txt"
                parts = [
                    f"IP: {ip}",
                    f"Collected at: {now()}",
                    f"identity: {name or '-'}",
                    f"device_type: {device_type or '-'}",
                    "",
                    "----- interfaces -----",
                    interfaces or "(empty)",
                    "",
                    "----- ip_addresses -----",
                    ip_addrs or "(empty)",
                    "",
                    "----- wireless_info -----",
                    wireless_info or "(empty)",
                    "",
                ]
                try:
                    prev = db.query_devices(q=ip, limit=1)
                    for d in prev.get("devices") or []:
                        if d.get("ip") == ip and d.get("ip_neighbors"):
                            parts.append("----- ip_neighbors -----")
                            parts.append(d["ip_neighbors"])
                            break
                except Exception:
                    pass
                fpath.write_text(chr(10).join(parts), encoding="utf-8")
            except Exception:
                pass
    return {
        "results": results,
        "saved": saved,
        "database": active_db,
        "txt_dir": str(txt_dir),
        "canceled": canceled,
        "probed": len(results),
        "requested": len(body.ips),
    }


@app.post("/api/snmp/probe/cancel")
def snmp_probe_cancel():
    """Request cancellation of the currently running SNMP probe."""
    PROBE_CANCEL.set()
    return {"cancel_requested": True, "message": "درخواست توقف خواندن اطلاعات ارسال شد"}


@app.get("/api/snmp/probe/progress")
def snmp_probe_progress():
    """Progress of the current (or last) SNMP probe."""
    p = dict(getattr(snmp_manager, "PROBE_PROGRESS", {}) or {})
    total = int(p.get("total") or 0)
    done = int(p.get("done") or 0)
    pct = int(round(100 * done / total)) if total else 0
    return {
        "total": total,
        "done": done,
        "percent": pct,
        "current_ip": p.get("current_ip") or "",
        "running": bool(p.get("running")),
        "canceled": bool(p.get("canceled")),
    }

class PushConfigRequest(BaseModel):
    ips: list[str] = Field(min_length=1, max_length=1000)
    username: str = Field(min_length=1, max_length=128)
    password: str = Field(default="", max_length=128)
    script: str = Field(min_length=1, max_length=50_000)
    port: int = Field(default=8728, ge=1, le=65535)
    timeout: float = Field(default=10.0, ge=1.0, le=60.0)



class ApiProbeRequest(BaseModel):
    ips: list[str] = Field(min_length=1, max_length=1000)
    username: str = Field(min_length=1, max_length=128)
    password: str = Field(default="", max_length=128)
    port: int = Field(default=8728, ge=1, le=65535)
    timeout: float = Field(default=12.0, ge=2.0, le=60.0)
    save: bool = True


@app.post("/api/devices/probe-api")
def probe_devices_api(body: ApiProbeRequest):
    """Read identity/interfaces/wireless mode/IP via RouterOS API and update DB."""
    API_PROBE_CANCEL.clear()
    API_PROBE_PROGRESS.update({"total": 0, "done": 0, "current_ip": "", "running": True, "canceled": False})
    results = routeros_api.probe_devices_via_api_many(
        body.ips,
        username=body.username,
        password=body.password or "",
        port=body.port,
        timeout=body.timeout,
        cancel_event=API_PROBE_CANCEL,
        progress=API_PROBE_PROGRESS,
    )
    canceled = API_PROBE_CANCEL.is_set()
    saved = 0
    active_db = db.get_active_db_name()
    from pathlib import Path as _P
    txt_dir = _P(db.db_path_for()).parent / "ip_info"
    txt_dir.mkdir(parents=True, exist_ok=True)
    if True:  # always persist API probe results to DB
        for r in results:
            if not r.get("ok"):
                continue
            ip = r["ip"]
            try:
                kwargs = {}
                if r.get("device_name"):
                    kwargs["device_name"] = r["device_name"]
                # Always persist detail fields when probe succeeded (even short text)
                if (r.get("interfaces") or "").strip():
                    kwargs["interfaces"] = r["interfaces"]
                if (r.get("ip_addresses") or "").strip():
                    kwargs["ip_addresses"] = r["ip_addresses"]
                if (r.get("ip_neighbors") or "").strip():
                    kwargs["ip_neighbors"] = r["ip_neighbors"]
                if (r.get("wireless_info") or "").strip():
                    kwargs["wireless_info"] = r["wireless_info"]
                if r.get("device_type"):
                    kwargs["device_type"] = r["device_type"]
                if (r.get("device_log") or "").strip():
                    kwargs["device_log"] = r["device_log"]
                if (r.get("mac_info") or "").strip():
                    kwargs["mac_info"] = r["mac_info"]
                if not kwargs:
                    continue
                try:
                    db.update_device_fields(ip, **kwargs)
                except KeyError:
                    db.add_manual_device(ip, open_ports=[8728, 8291], classification="mikrotik_candidate")
                    db.update_device_fields(ip, **kwargs)
                saved += 1
                safe_ip = ip.replace(":", "_").replace("/", "_")
                fpath = txt_dir / f"{safe_ip}.txt"
                parts = [
                    f"IP: {ip}",
                    f"Collected at: {now()}",
                    f"source: RouterOS API",
                    f"identity: {r.get('device_name') or '-'}",
                    f"device_type: {r.get('device_type') or '-'}",
                    "",
                    "----- interfaces -----",
                    r.get("interfaces") or "(empty)",
                    "",
                    "----- wireless -----",
                    r.get("wireless_info") or "(empty)",
                    "",
                    "----- ip_addresses / routes -----",
                    r.get("ip_addresses") or "(empty)",
                    "",
                    "----- neighbors -----",
                    r.get("ip_neighbors") or "(empty)",
                    "",
                    "----- log -----",
                    r.get("device_log") or "(empty)",
                    "",
                    "----- mac -----",
                    r.get("mac_info") or "(empty)",
                ]
                fpath.write_text(chr(10).join(parts), encoding="utf-8")
            except Exception:
                pass
    ok = sum(1 for r in results if r.get("ok"))
    return {
        "results": results,
        "ok": ok,
        "failed": len(results) - ok,
        "saved": saved,
        "canceled": canceled,
        "probed": len(results),
        "requested": len(body.ips),
        "database": active_db,
        "txt_dir": str(txt_dir),
    }


@app.post("/api/devices/probe-api/cancel")
def probe_devices_api_cancel():
    API_PROBE_CANCEL.set()
    return {"cancel_requested": True, "message": "درخواست توقف خواندن API ارسال شد"}


@app.get("/api/devices/probe-api/progress")
def probe_devices_api_progress():
    p = dict(API_PROBE_PROGRESS or {})
    total = int(p.get("total") or 0)
    done = int(p.get("done") or 0)
    pct = int(round(100 * done / total)) if total else 0
    return {
        "total": total,
        "done": done,
        "percent": pct,
        "current_ip": p.get("current_ip") or "",
        "running": bool(p.get("running")),
        "canceled": bool(p.get("canceled")),
    }


@app.post("/api/devices/push-config")
def push_config_to_devices(body: PushConfigRequest):
    """Connect to selected devices via RouterOS API (8728) and run a terminal script."""
    script = (body.script or "").strip()
    if not script:
        raise HTTPException(400, "متن کانفیگ خالی است")
    ips = [str(x).strip() for x in body.ips if str(x).strip()]
    if not ips:
        raise HTTPException(400, "هیچ IP معتبری ارسال نشده")
    results = routeros_api.push_config_many(
        ips,
        username=body.username,
        password=body.password or "",
        script=script,
        port=body.port,
        timeout=body.timeout,
    )
    ok = sum(1 for r in results if r.get("ok"))
    return {
        "results": results,
        "ok": ok,
        "failed": len(results) - ok,
        "total": len(results),
    }


class DeviceMetaRequest(BaseModel):
    device_name: str | None = Field(default=None, max_length=128)
    tags: str | None = Field(default=None, max_length=4000)
    ip_addresses: str | None = Field(default=None, max_length=100000)
    ip_neighbors: str | None = Field(default=None, max_length=100000)
    interfaces: str | None = Field(default=None, max_length=100000)
    device_type: str | None = Field(default=None, max_length=32)
    wireless_info: str | None = Field(default=None, max_length=100000)
    device_log: str | None = Field(default=None, max_length=200000)
    mac_info: str | None = Field(default=None, max_length=100000)


@app.patch("/devices/{ip}")
def patch_device(ip: str, body: DeviceMetaRequest):
    try:
        return db.update_device_fields(
            ip,
            device_name=body.device_name,
            tags=body.tags,
            ip_addresses=body.ip_addresses,
            ip_neighbors=body.ip_neighbors,
            interfaces=body.interfaces,
            device_type=body.device_type,
            wireless_info=body.wireless_info,
            device_log=body.device_log,
            mac_info=body.mac_info,
        )
    except KeyError:
        raise HTTPException(404, "دستگاه یافت نشد") from None
    except ValueError as e:
        raise HTTPException(422, str(e)) from e


class CollectIpInfoRequest(BaseModel):
    ips: list[str] = Field(min_length=1, max_length=1000)
    username: str = Field(min_length=1, max_length=128)
    password: str = Field(default="", max_length=128)
    port: int = Field(default=8728, ge=1, le=65535)
    timeout: float = Field(default=12.0, ge=1.0, le=60.0)
    save_txt: bool = Field(default=True)


@app.post("/api/devices/collect-ip-info")
def collect_ip_info_endpoint(body: CollectIpInfoRequest):
    """Collect /ip address + /ip route into ip_addresses field and /ip neighbor into ip_neighbors.
    Optionally write each result into a txt file under databases/ip_info/."""
    ips = [str(x).strip() for x in body.ips if str(x).strip()]
    if not ips:
        raise HTTPException(400, "هیچ IP معتبری ارسال نشده")
    results = routeros_api.collect_ip_info_many(
        ips,
        username=body.username,
        password=body.password or "",
        port=body.port,
        timeout=body.timeout,
    )
    saved = 0
    txt_dir = None
    if body.save_txt:
        from pathlib import Path as _P
        txt_dir = _P(db.db_path_for()).parent / "ip_info"
        txt_dir.mkdir(parents=True, exist_ok=True)
    for r in results:
        if not r.get("ok"):
            continue
        try:
            db.update_device_fields(
                r["ip"],
                ip_addresses=r.get("ip_addresses") or "",
                ip_neighbors=r.get("ip_neighbors") or "",
            )
            saved += 1
            if txt_dir is not None:
                safe_ip = r["ip"].replace(":", "_").replace("/", "_")
                fpath = txt_dir / f"{safe_ip}.txt"
                text = (
                    f"IP: {r['ip']}\n"
                    f"Collected at: {now()}\n\n"
                    f"{r.get('ip_addresses') or ''}\n\n"
                    f"{r.get('ip_neighbors') or ''}\n"
                )
                fpath.write_text(text, encoding="utf-8")
        except Exception:
            pass
    ok = sum(1 for r in results if r.get("ok"))
    return {
        "results": results,
        "ok": ok,
        "failed": len(results) - ok,
        "saved": saved,
        "txt_dir": str(txt_dir) if txt_dir else None,
        "database": db.get_active_db_name(),
    }


# Run with one of:
#   python backend/main.py      (from project root)
#   python main.py              (from inside backend/)
#   python -m backend.main      (from project root)

# ---- Serve embedded UI (same-origin at http://127.0.0.1:8001/) ----
# Avoids Chrome Private Network Access block when UI is on a public host.
try:
    from fastapi.staticfiles import StaticFiles
    from fastapi.responses import FileResponse
    _UI_CANDIDATES = [
        ROOT_DIR / "frontend" / "dist",
        ROOT_DIR / "ui",
        Path(__file__).resolve().parent.parent / "ui",
        Path(__file__).resolve().parent / "ui",
    ]
    if getattr(__import__("sys"), "frozen", False) and hasattr(__import__("sys"), "_MEIPASS"):
        _UI_CANDIDATES.insert(0, Path(__import__("sys")._MEIPASS) / "ui")
    _UI_DIR = next((d for d in _UI_CANDIDATES if d.is_dir() and (d / "index.html").exists()), None)
    if _UI_DIR is not None:
        @app.get("/")
        def _ui_index():
            return FileResponse(_UI_DIR / "index.html")
        app.mount("/assets", StaticFiles(directory=str(_UI_DIR / "assets")), name="ui-assets")
        print(f"[ui] serving frontend from {_UI_DIR}")
    else:
        print("[ui] no frontend dist found — API only")
except Exception as _ui_err:
    print(f"[ui] static mount skipped: {_ui_err}")

if __name__ == "__main__":
    try:
        import uvicorn
    except ImportError:
        raise SystemExit(
            "uvicorn is not installed. Run:  pip install -r requirements.txt  "
            "(or pip install uvicorn[standard])"
        )
    # app_dir guarantees uvicorn can import "backend.main" even when the
    # current working directory is backend/ itself; reload=True still works.
    uvicorn.run(
        "backend.main:app",
        host="0.0.0.0",
        port=8001,
        reload=True,
        app_dir=str(ROOT_DIR),
    )
