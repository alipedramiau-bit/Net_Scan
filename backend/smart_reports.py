"""Interface link flap detection and Excel report helpers."""
from __future__ import annotations
from datetime import datetime, timedelta
import re
from io import BytesIO, StringIO
import csv

_LINE = re.compile(
    r"^\s*(?:\d+\s+)?(?P<stamp>(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|"
    r"[A-Za-z]{3}/\d{1,2})\s+\d{2}:\d{2}:\d{2}|"
    r"\d{2}:\d{2}:\d{2})\s+(?P<message>.*)$", re.I)
# Topics may be single ("info") or comma-separated ("interface,info,debug").
# Interface names may include /, ., @, :, <>, e.g. sfp-sfpplus1, <pppoe-out1>, ether1.
_LINK_EVENT = re.compile(
    r"^\s*(?:\[[^]]*\]\s*)?"
    r"(?:(?P<topics>[\w-]+(?:,[\w-]+)*)\s+)?"
    r"(?P<interface>[\w./@<>:+-]+?)\s*:?\s+link\s+(?P<status>up|down)\b",
    re.I,
)
_MONTHS = {name.lower(): number for number, name in enumerate(
    ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"), 1)}
_GREGORIAN_YEAR = 1600

def jalali_to_gregorian(jy, jm, jd):
    """Convert Jalali calendar date to Gregorian (integer arithmetic)."""
    jy += 1595
    days = -355668 + 365 * jy + (jy // 33) * 8 + ((jy % 33 + 3) // 4) + jd
    if jm < 7: days += (jm - 1) * 31
    else: days += (jm - 7) * 30 + 186
    gy = 400 * (days // 146097)
    days %= 146097
    if days > 36524:
        gy += 100 * ((days - 1) // 36524); days = (days - 1) % 36524
        if days >= 365: days += 1
    gy += 4 * (days // 1461); days %= 1461
    if days > 365:
        gy += (days - 1) // 365; days = (days - 1) % 365
    gd = days + 1
    leap = (gy % 4 == 0 and gy % 100 != 0) or gy % 400 == 0
    month_days = [0,31,29 if leap else 28,31,30,31,30,31,31,30,31,30,31]
    gm = 1
    while gm <= 12 and gd > month_days[gm]:
        gd -= month_days[gm]; gm += 1
    return gy, gm, gd

def gregorian_to_jalali(gy, gm, gd):
    """Convert Gregorian date to Jalali using integer day arithmetic."""
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    j_days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    gy -= 1600
    g_day_no = (365 * gy + (gy + 3) // 4 - (gy + 99) // 100 + (gy + 399) // 400)
    for month in range(1, gm):
        g_day_no += g_days_in_month[month - 1]
    if gm > 2 and ((gy % 4 == 0 and gy % 100 != 0) or gy % 400 == 0):
        g_day_no += 1
    g_day_no += gd - 1

    j_day_no = g_day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053
    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461
    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365
    jm = 1
    while jm <= 12 and j_day_no >= j_days_in_month[jm - 1]:
        j_day_no -= j_days_in_month[jm - 1]
        jm += 1
    return jy, jm, j_day_no + 1


def jalali_datetime_string(value):
    """Format a datetime as Jalali YYYY/MM/DD HH:MM:SS."""
    jy, jm, jd = gregorian_to_jalali(value.year, value.month, value.day)
    return f"{jy:04d}/{jm:02d}/{jd:02d} {value:%H:%M:%S}"


def parse_log_datetime(value, reference=None):
    """Parse Gregorian, legacy Jalali, abbreviated-month, and time-only dates.

    Yearless dates are assigned the year nearest the supplied reference (or
    today), which naturally handles logs spanning a year boundary.
    """
    reference = reference or datetime.now()
    value = value.strip()
    if re.fullmatch(r"\d{2}:\d{2}:\d{2}", value):
        return datetime.combine(reference.date(), datetime.strptime(value, "%H:%M:%S").time())
    month_match = re.fullmatch(r"([A-Za-z]{3})/(\d{1,2})\s+(\d{2}:\d{2}:\d{2})", value)
    if month_match:
        month = _MONTHS.get(month_match.group(1).lower())
        if not month:
            raise ValueError("Unknown month abbreviation")
        day, clock = int(month_match.group(2)), datetime.strptime(month_match.group(3), "%H:%M:%S").time()
        candidates = []
        for year in (reference.year - 1, reference.year, reference.year + 1):
            try:
                candidates.append(datetime(year, month, day, clock.hour, clock.minute, clock.second))
            except ValueError:
                pass
        if not candidates:
            raise ValueError("Invalid yearless date")
        return min(candidates, key=lambda item: abs((item - reference).total_seconds()))
    normalized = value.replace("-", "/")
    dt = datetime.strptime(normalized, "%Y/%m/%d %H:%M:%S")
    if dt.year < _GREGORIAN_YEAR:
        gy, gm, gd = jalali_to_gregorian(dt.year, dt.month, dt.day)
        return dt.replace(year=gy, month=gm, day=gd)
    return dt

def parse_device_log(device, now=None, days_back=5, threshold_minutes=5):
    """Return latest flap per interface.

    In addition to the existing short-interval detection, report an interface
    from the third event onward in a consecutive alternating sequence (at
    least two transitions), provided each adjacent event is less than one hour
    apart.
    """
    now = now or datetime.now()
    cutoff = now - timedelta(days=max(0, int(days_back)))
    grouped = {}
    for line in (device.get("device_log") or "").splitlines():
        match = _LINE.match(line)
        if not match: continue
        event = _LINK_EVENT.match(match.group("message"))
        if not event: continue
        try: when = parse_log_datetime(match.group("stamp"), reference=now)
        except (ValueError, OverflowError): continue
        if when < cutoff or when > now + timedelta(minutes=1): continue
        iface = event.group("interface").strip().rstrip(":").strip()
        # Drop accidental topic words captured as interface when topics group missed.
        if iface.lower() in {"info", "debug", "warning", "error", "critical", "system", "interface", "link", "wireless"}:
            continue
        status = event.group("status").title()
        if iface:
            grouped.setdefault(iface, []).append((when, status))
    found=[]
    threshold=timedelta(minutes=max(0, int(threshold_minutes)))
    one_hour=timedelta(hours=1)
    for iface, events in grouped.items():
        events.sort(key=lambda item:item[0])
        flaps=[]
        alternating_transitions=0
        for prev, cur in zip(events, events[1:]):
            gap=cur[0]-prev[0]
            if prev[1] != cur[1]:
                if timedelta(0) <= gap <= threshold:
                    flaps.append(cur)  # Preserve the existing short-flap rule.
                if timedelta(0) <= gap < one_hour:
                    alternating_transitions += 1
                    if alternating_transitions >= 2:
                        flaps.append(cur)
                else:
                    alternating_transitions=0
            else:
                alternating_transitions=0
        # A record can satisfy both rules; count unique event records once.
        flaps=list(dict.fromkeys(flaps))
        if flaps:
            latest=max(flaps,key=lambda item:item[0])
            raw_tags = device.get('tags') or device.get('device_tag') or device.get('tag') or ''
            if isinstance(raw_tags, list):
                tag_str = ', '.join(str(t).strip() for t in raw_tags if str(t).strip())
            else:
                tag_str = str(raw_tags).strip()
            found.append({"datetime":latest[0].isoformat(sep=" "),
                          "datetime_jalali":jalali_datetime_string(latest[0]),
                          "device_name":device.get("device_name") or "",
                          "tags":tag_str, "device_tag": tag_str,
                          "ip":device.get("ip") or "", "interface":iface, "last_status":latest[1],
                          "flap_count":len(flaps)})
    return found

def interface_flap_report(devices, days_back=5, threshold_minutes=5, now=None):
    now = now or datetime.now()
    devices=list(devices or [])
    events=[]
    devices_with_logs=0
    newest_log=None
    for device in devices:
        log_text=device.get("device_log") or ""
        if log_text.strip(): devices_with_logs+=1
        # Track the newest link event in the logs regardless of the requested
        # window, so an empty result can be explained (stale logs vs. none).
        for line in log_text.splitlines():
            match=_LINE.match(line)
            if not match: continue
            if not _LINK_EVENT.match(match.group("message")): continue
            try: when=parse_log_datetime(match.group("stamp"), reference=now)
            except (ValueError, OverflowError): continue
            if newest_log is None or when>newest_log: newest_log=when
        events.extend(parse_device_log(device, now=now, days_back=days_back, threshold_minutes=threshold_minutes))
    events.sort(key=lambda e:(e["datetime"],e["device_name"],e["interface"]),reverse=True)
    for index,event in enumerate(events,1): event["row"] = index
    report={"events":events,"device_count":len({e["ip"] for e in events}),"event_count":len(events),
            "days":int(days_back),"threshold_minutes":int(threshold_minutes),
            "devices_scanned":len(devices),"devices_with_logs":devices_with_logs}
    if newest_log is not None:
        report["newest_log_at"]=newest_log.isoformat(sep=" ")
        report["newest_log_at_jalali"]=jalali_datetime_string(newest_log)
        if not events:
            cutoff=now-timedelta(days=max(0,int(days_back)))
            if newest_log<cutoff:
                report["stale_logs"]=True
                suggested=max(int(days_back)+1, min(90, int((now-newest_log).days)+2))
                report["suggested_days"]=suggested
    return report

def interface_flaps_csv(report):
    """Return a UTF-8-BOM CSV for Excel, populated from a flap report.

    The CSV mirrors every field shown in the on-screen report table,
    using the same column names and order as the UI.
    """
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["ردیف", "نام دستگاه", "آی‌پی دستگاه", "تگ دستگاه", "اینترفیس", "تعداد تکرار لاگ", "تاریخ و زمان (میلادی)", "تاریخ و زمان (شمسی)", "نوع رویداد"])
    for event in report.get("events", []):
        writer.writerow([event.get("row", ""), event.get("device_name", ""),
                         event.get("ip", ""), event.get("device_tag", event.get("tags", "")),
                         event.get("interface", ""), event.get("flap_count", 1),
                         event.get("datetime", ""), event.get("datetime_jalali", ""),
                         event.get("last_status", "")])
    return "\ufeff" + output.getvalue()

def interface_flaps_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    wb=Workbook(); ws=wb.active; ws.title="قطع وصل اینترفیس"; ws.sheet_view.rightToLeft=True
    headers=["ردیف","نام دستگاه","آی‌پی دستگاه","تگ دستگاه","اینترفیس","تعداد تکرار لاگ","تاریخ و زمان (میلادی)","تاریخ و زمان (شمسی)","نوع رویداد"]
    ws.append(headers)
    for event in report.get("events",[]):
        ws.append([event["row"],event["device_name"],event["ip"],event.get("device_tag") or event.get("tags") or "",event["interface"],event.get("flap_count",1),event["datetime"],event["datetime_jalali"],event["last_status"]])
    fill=PatternFill("solid",fgColor="1F4E78"); white=Font(name="Arial",bold=True,color="FFFFFF",size=11)
    thin=Side(style="thin",color="D9E2F3")
    for cell in ws[1]:
        cell.fill=fill;cell.font=white;cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True);cell.border=Border(bottom=Side(style="medium",color="17365D"))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name="Arial",size=10);cell.alignment=Alignment(horizontal="center",vertical="center");cell.border=Border(bottom=thin)
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=32
    for col in ws.columns:
        letter=col[0].column_letter; width=max((len(str(c.value or "")) for c in col),default=10)
        ws.column_dimensions[letter].width=min(max(width+3,12),42)
    out=BytesIO();wb.save(out);out.seek(0);return out


# ---------------- Smart report: wireless antenna / station disconnects ----------------
_MAC_ADDRESS = re.compile(r"(?i)\b(?:[0-9a-f]{2}[:-]){5}[0-9a-f]{2}\b")
_WIRELESS_ISSUE = re.compile(
    r"(?:"
    r"extensive\s+data\s+loss|"
    r"lost\s+connection|connection\s+lost|"
    r"disassoc(?:iated|iation)?|"
    r"deauth(?:enticated|entication)?|"
    r"station\s+leaving|"
    r"not\s+responding|no\s+response|"
    r"no\s+beacons?\s+received|"
    r"failed\s+to\s+connect|"
    r"registration\s+timeout|authentication\s+timeout|"
    r"removing\s+from\s+registration\s+table|"
    r"(?:weak|low|poor)\s+(?:wireless\s+)?signal|"
    r"signal\s+(?:is\s+)?(?:weak|low|lost|too\s+weak|below)"
    r")",
    re.I,
)

def _mac_key(value):
    return re.sub(r"[^0-9a-f]", "", str(value or "").lower())

def wireless_antenna_report(devices, days_back=5, now=None):
    """Find wireless station/link issues in saved RouterOS logs and resolve MACs against DB fields."""
    now = now or datetime.now()
    cutoff = now - timedelta(days=max(1, int(days_back)))
    devices = list(devices or [])
    # Index MACs found in saved inventory fields to their known device identity/IP.
    mac_index = {}
    mac_fields = ("mac_info", "wireless_info", "ip_neighbors", "interfaces", "ip_addresses")
    for device in devices:
        for field in mac_fields:
            text_value = device.get(field) or ""
            for mac in _MAC_ADDRESS.findall(text_value):
                key = _mac_key(mac)
                if key and key not in mac_index:
                    mac_index[key] = {
                        "matched_device_ip": device.get("ip") or "",
                        "matched_device_name": device.get("device_name") or "",
                    }

    events = []
    seen = set()
    for device in devices:
        raw_log = device.get("device_log") or ""
        for line in raw_log.splitlines():
            match = _LINE.match(line)
            if not match:
                continue
            message = match.group("message").strip()
            if not _WIRELESS_ISSUE.search(message):
                continue
            try:
                when = parse_log_datetime(match.group("stamp"), reference=now)
            except (ValueError, OverflowError):
                continue
            if when < cutoff or when > now + timedelta(minutes=1):
                continue
            macs = _MAC_ADDRESS.findall(message)
            # Signal-strength-only records without a station MAC are still useful,
            # but the report cannot claim a client identity without one.
            station_mac = macs[0].upper().replace("-", ":") if macs else ""
            resolved = mac_index.get(_mac_key(station_mac), {}) if station_mac else {}
            key = (device.get("ip") or "", when.isoformat(), message)
            if key in seen:
                continue
            seen.add(key)
            events.append({
                "datetime": when.isoformat(sep=" "),
                "datetime_jalali": jalali_datetime_string(when),
                "antenna_ip": device.get("ip") or "",
                "antenna_name": device.get("device_name") or "",
                "station_mac": station_mac,
                "matched_device_ip": resolved.get("matched_device_ip", ""),
                "matched_device_name": resolved.get("matched_device_name", ""),
                "event": message,
                "raw_line": line.strip(),
            })
    events.sort(key=lambda item: (item["datetime"], item["antenna_ip"], item["station_mac"]), reverse=True)
    for index, event in enumerate(events, 1):
        event["row"] = index
    return {
        "items": events,
        "events": events,
        "event_count": len(events),
        "days": int(days_back),
        "unmatched_mac_count": len({e["station_mac"] for e in events if e["station_mac"] and not e["matched_device_ip"]}),
    }

def wireless_antenna_csv(report):
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["ردیف", "تاریخ و زمان (میلادی)", "تاریخ و زمان (شمسی)",
                     "نام آنتن/روتر لاگ‌دهنده", "IP آنتن/روتر", "MAC ایستگاه",
                     "نام دستگاه متناظر با MAC", "IP دستگاه متناظر با MAC", "رویداد/علت"])
    for event in report.get("items", []):
        writer.writerow([
            event.get("row", ""), event.get("datetime", ""), event.get("datetime_jalali", ""),
            event.get("antenna_name", ""), event.get("antenna_ip", ""), event.get("station_mac", ""),
            event.get("matched_device_name", ""), event.get("matched_device_ip", ""), event.get("event", ""),
        ])
    return "\ufeff" + output.getvalue()
